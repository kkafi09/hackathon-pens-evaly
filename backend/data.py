"""Read-only source index. Every graph edge is backed by a source record."""
import csv
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

SOURCE = Path(__file__).resolve().parents[1] / 'data' / 'dataset_kasirnusa'
SNAPSHOT = '2026-10-01'
IDS = {'crm_accounts': 'account_id', 'crm_contacts': 'contact_id', 'employees': 'employee_id',
       'contracts_billing': 'contract_id', 'crm_deals': 'deal_id', 'decision_log': 'decision_id',
       'support_tickets': 'ticket_id', 'bugs': 'bug_id', 'features': 'feature_id', 'outlets': 'outlet_id',
       'interactions': 'interaction_id', 'releases': 'versi'}


class Dataset:
    def __init__(self, source=SOURCE):
        self.source = Path(source)
        self.tables = {}
        digest = hashlib.sha256()
        for file in sorted(self.source.iterdir()):
            if file.suffix in {'.csv', '.jsonl', '.xlsx'}:
                digest.update(file.name.encode())
                with file.open('rb') as stream:
                    for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                        digest.update(chunk)
            if file.suffix == '.csv':
                with file.open(encoding='utf-8-sig', newline='') as stream:
                    self.tables[file.stem] = list(csv.DictReader(stream))
            elif file.suffix == '.jsonl':
                self.tables[file.stem] = [json.loads(line) for line in file.read_text(encoding='utf-8-sig').splitlines() if line.strip()]
        self.hash = digest.hexdigest()
        self.warnings = []
        self._reconcile_excel()
        self._prepare()

    @classmethod
    def from_records(cls, tables, digest, warnings):
        instance = cls.__new__(cls)
        instance.tables = tables
        instance.hash = digest
        instance.warnings = warnings
        instance._prepare()
        return instance

    def _prepare(self):
        self.index = {}
        for table, key in IDS.items():
            records = self.tables[table]
            self.index[table] = {r[key]: r for r in records}
            if len(self.index[table]) != len(records):
                raise ValueError(f'ID duplikat di {table}')
        self.accounts = self.index['crm_accounts']
        self.usage = defaultdict(lambda: defaultdict(lambda: {'records': 0, 'transactions': 0, 'outlets': set(), 'versions': set(), 'offline_records': 0}))
        for r in self.tables.pop('product_usage_daily'):
            month = r['tanggal'][:7]
            item = self.usage[r['account_id']][month]
            item['records'] += 1
            item['transactions'] += int(r['jumlah_transaksi'])
            item['outlets'].add(r['outlet_id'])
            item['versions'].add(r['versi_aplikasi'])
            item['offline_records'] += int(bool(r['transaksi_offline_tersinkron']))
        self._validate_links()

    def _reconcile_excel(self):
        from openpyxl import load_workbook
        book = load_workbook(self.source / 'decision_log.xlsx', read_only=True, data_only=True)
        excel = {}
        for sheet in book:
            values = iter(sheet.values)
            headers = next(values, ())
            if 'decision_id' not in headers:
                continue
            for values_row in values:
                row = dict(zip(headers, values_row))
                if row.get('decision_id'):
                    excel[row['decision_id']] = row
        book.close()
        csv_rows = {row['decision_id']: row for row in self.tables['decision_log']}
        for key in sorted(set(excel) | set(csv_rows)):
            if key not in excel or key not in csv_rows:
                self.warnings.append(f'{key}: keputusan tidak ada pada salah satu CSV/Excel')
                continue
            for col, value in csv_rows[key].items():
                other = excel[key].get(col)
                if isinstance(other, datetime):
                    other = other.date().isoformat()
                if str(other if other is not None else '').strip() != value.strip():
                    try:
                        if col == 'nilai' and value.strip().endswith('%'):
                            same = Decimal(str(other)) == Decimal(value.strip()[:-1]) / 100
                        else:
                            same = Decimal(str(other or 0)) == Decimal(value or 0)
                    except (ValueError, InvalidOperation):
                        same = False
                    if not same:
                        self.warnings.append(f'{key}.{col}: CSV dan Excel berbeda; CSV digunakan, nilai Excel tidak digabung.')

    def _validate_links(self):
        for table, rows in self.tables.items():
            for r in rows:
                aid = r.get('account_id')
                if aid and aid not in self.accounts:
                    self.warnings.append(f'{table}: akun {aid} tidak ditemukan')
                outlet = r.get('outlet_id')
                if table == 'support_tickets' and outlet:
                    linked = self.index['outlets'].get(outlet)
                    if not linked or linked['account_id'] != aid:
                        self.warnings.append(f"{r['ticket_id']}: relasi outlet tidak cocok")

    def context(self, aid):
        account = self.accounts[aid]
        history = [r for r in self.tables['contact_employment_history'] if r['account_id'] == aid]
        contact_ids = {r['contact_id'] for r in history}
        contacts = [r for r in self.tables['crm_contacts'] if r['account_id_saat_ini'] == aid or r['contact_id'] in contact_ids]
        contacts_by_email = {r['email'].lower(): r for r in contacts}
        # An old employer domain can differ. Match a historical address only when
        # its local part identifies exactly one known contact, then verify tenure.
        local_parts = defaultdict(list)
        for r in self.tables['crm_contacts']:
            local_parts[r['email'].lower().split('@')[0]].append(r)
        def resolve_sender(email):
            exact = contacts_by_email.get(email.lower())
            candidates = local_parts.get(email.lower().split('@')[0], [])
            return exact or (candidates[0] if len(candidates) == 1 and candidates[0]['contact_id'] in contact_ids else None)
        direct = {r['interaction_id'] for r in self.tables['interactions'] if r['account_id'] == aid}
        for r in self.tables['interactions']:
            if r['account_id']:
                continue
            addresses = re.findall(r'[\w.+-]+@[\w.-]+', r['dari'] + ';' + r['ke'])
            participants = set(r['peserta'].split(';'))
            for c in contacts:
                addressed = c['email'].lower() in [a.lower() for a in addresses] or c['contact_id'] in participants
                valid = any(h['contact_id'] == c['contact_id'] and h['mulai'] <= r['tanggal'] and (not h['selesai'] or r['tanggal'] <= h['selesai']) for h in history)
                if addressed and valid:
                    direct.add(r['interaction_id'])
        # Follow replies in both directions, excluding records explicitly belonging to another account.
        changed = True
        while changed:
            changed = False
            for r in self.tables['interactions']:
                if r['account_id'] not in ('', aid):
                    continue
                if r['interaction_id'] in direct and r['membalas_id']:
                    parent = self.index['interactions'].get(r['membalas_id'])
                    if parent and parent['account_id'] in ('', aid) and parent['interaction_id'] not in direct:
                        direct.add(parent['interaction_id']); changed = True
                elif r['membalas_id'] in direct and r['interaction_id'] not in direct:
                    direct.add(r['interaction_id']); changed = True
        interactions = sorted([dict(r) for r in self.tables['interactions'] if r['interaction_id'] in direct], key=lambda r: r['tanggal'])
        for r in interactions:
            r['source'] = 'interactions.jsonl'
            sender = resolve_sender(r['dari'])
            r['customer_sender'] = bool(sender and any(h['contact_id'] == sender['contact_id'] and h['mulai'] <= r['tanggal'] and (not h['selesai'] or r['tanggal'] <= h['selesai']) for h in history))
            r['sender_contact_id'] = sender['contact_id'] if r['customer_sender'] else None
            r['reply_count'] = sum(x['membalas_id'] == r['interaction_id'] for x in interactions)
        def own(table):
            return [r for r in self.tables[table] if r.get('account_id') == aid]
        tickets = own('support_tickets')
        bugs = [self.index['bugs'][key] for key in sorted({r['bug_id'] for r in tickets if r['bug_id'] in self.index['bugs']})]
        monthly = [{'month': month, **{k: sorted(v) if isinstance(v, set) else v for k, v in values.items()},
                    'transactions_per_outlet_day': round(values['transactions'] / max(values['records'], 1), 2)}
                   for month, values in sorted(self.usage[aid].items())]
        decisions = own('decision_log')
        precedents = [r for r in self.tables['decision_log'] if r['account_id'] != aid and r['tipe'] in ('diskon', 'pengecualian', 'janji_fitur')]
        result = {'account': account, 'owner': self.index['employees'].get(account['account_owner_id']),
                  'contacts': contacts, 'employment_history': history, 'interactions': interactions,
                  'contracts': own('contracts_billing'), 'deals': own('crm_deals'), 'decisions': decisions,
                  'tickets': tickets, 'outlets': own('outlets'), 'usage': monthly,
                  'feature_usage': own('feature_usage_monthly'), 'features': self.tables['features'],
                  'bugs': bugs, 'releases': self.tables['releases'], 'precedents': precedents,
                  'snapshot': SNAPSHOT, 'dataset_hash': self.hash, 'warnings': self.warnings}
        result['source_counts'] = {key: len(result[key]) for key in ('contacts', 'employment_history', 'interactions', 'contracts', 'deals', 'decisions', 'tickets', 'outlets', 'feature_usage')}
        result['source_counts']['product_usage_daily'] = sum(x['records'] for x in monthly)
        return result

    def graph(self, aid=None):
        nodes, edges = {}, {}
        def node(key, kind, label, record, source):
            nodes[key] = {'id': key, 'kind': kind, 'label': label, 'record': record, 'source': source}
        def edge(source, target, label, evidence):
            if source not in nodes or target not in nodes:
                return
            key = f'{source}:{label}:{target}'
            edges[key] = {'id': key, 'source': source, 'target': target, 'label': label, 'evidence': evidence, 'inference': False}
        def employee(eid):
            r = self.index['employees'].get(eid)
            if r:
                node(eid, 'employee', r['nama'], r, 'employees.csv')
        accounts = [self.accounts[aid]] if aid else self.accounts.values()
        for r in accounts:
            key = r['account_id']
            node(key, 'account', r['nama'], r, 'crm_accounts.csv')
            employee(r['account_owner_id'])
            edge(key, r['account_owner_id'], 'DIPEGANG_OLEH', key)
        if aid:
            c = self.context(aid)
            for table, key, kind, label in [('contacts', 'contact_id', 'contact', 'nama'), ('interactions', 'interaction_id', 'interaction', 'subjek'),
                                          ('contracts', 'contract_id', 'contract', 'paket'), ('tickets', 'ticket_id', 'ticket', 'judul'),
                                          ('deals', 'deal_id', 'deal', 'stage'), ('decisions', 'decision_id', 'decision', 'tipe'),
                                          ('bugs', 'bug_id', 'bug', 'judul')]:
                source = {'contacts': 'crm_contacts.csv', 'interactions': 'interactions.jsonl', 'contracts': 'contracts_billing.csv',
                          'tickets': 'support_tickets.csv', 'deals': 'crm_deals.csv', 'decisions': 'decision_log.csv', 'bugs': 'bugs.csv'}[table]
                for r in c[table]:
                    node(r[key], kind, r[label], r, source)
                    if kind != 'bug':
                        edge(aid, r[key], 'TERKAIT' if kind in ('contact', 'interaction') else 'MEMILIKI', r[key])
            for h in c['employment_history']:
                edge(h['contact_id'], aid, 'PERNAH_BEKERJA_DI' if h['selesai'] else 'BEKERJA_DI', h)
            edge(c['account']['champion_contact_id'], aid, 'CHAMPION_DARI', aid)
            for r in c['interactions']:
                edge(r['interaction_id'], r['membalas_id'], 'MEMBALAS', r['interaction_id'])
                if r['sender_contact_id']:
                    edge(r['sender_contact_id'], r['interaction_id'], 'MENGIRIM', r['interaction_id'])
                for person in c['contacts']:
                    if person['contact_id'] in r['peserta'].split(';'):
                        edge(person['contact_id'], r['interaction_id'], 'TERLIBAT_DI', r['interaction_id'])
            for r in c['tickets']:
                edge(r['pelapor_contact_id'], r['ticket_id'], 'MEMBUKA_TIKET', r['ticket_id'])
                edge(r['ticket_id'], r['bug_id'], 'DISEBABKAN_OLEH', r['ticket_id'])
            used = {r['feature_id'] for r in c['feature_usage']} | {r['fitur_terkait'] for r in c['bugs']} | {r['fitur_dijanjikan'] for r in c['decisions']}
            for r in c['features']:
                if r['feature_id'] in used:
                    node(r['feature_id'], 'feature', r['nama'], r, 'features.csv')
                    if any(x['feature_id'] == r['feature_id'] for x in c['feature_usage']):
                        edge(aid, r['feature_id'], 'MENGGUNAKAN', 'feature_usage_monthly.csv')
            for r in c['bugs']:
                edge(r['bug_id'], r['fitur_terkait'], 'TERKAIT_FITUR', r['bug_id'])
            for r in c['decisions']:
                employee(r['diputuskan_oleh'])
                edge(r['diputuskan_oleh'], r['decision_id'], 'MENYETUJUI' if r['keputusan'] == 'Disetujui' else 'MEMUTUSKAN', r['decision_id'])
                edge(r['decision_id'], r['bukti_interaction_id'], 'DIDASARKAN_PADA', r['decision_id'])
                edge(r['decision_id'], r['fitur_dijanjikan'], 'MENJANJIKAN', r['decision_id'])
                edge(r['deal_id'], r['decision_id'], 'DIDASARKAN_PADA', r['decision_id'])
            for r in c['contracts']:
                edge(r['contract_id'], r['decision_id'], 'DIDASARKAN_PADA', r['contract_id'])
            versions = {v for m in c['usage'] for v in m['versions']}
            for r in c['releases']:
                if r['versi'] in versions:
                    rid = f"release:{r['versi']}"
                    node(rid, 'release', f"Rilis {r['versi']}", r, 'releases.csv')
                    edge(aid, rid, 'TERCATAT_PADA_VERSI', 'product_usage_daily.csv')
            uid = f'usage:{aid}'
            node(uid, 'usage', 'Penggunaan produk · 12 bulan', {'monthly': c['usage']}, 'product_usage_daily.csv')
            edge(aid, uid, 'MENGGUNAKAN', 'product_usage_daily.csv')
        pipeline = [r for r in self.tables['crm_deals'] if r['status'] == 'Terbuka' and self.accounts[r['account_id']]['tipe'] == 'prospek']
        return {'nodes': list(nodes.values()), 'edges': list(edges.values()), 'snapshot': SNAPSHOT, 'dataset_hash': self.hash,
                'pipeline': pipeline, 'pipeline_value': sum(int(r['nilai_tahunan']) for r in pipeline),
                'counts': {'customers': sum(r['tipe'] == 'pelanggan' for r in self.accounts.values()),
                           'prospects': sum(r['tipe'] == 'prospek' for r in self.accounts.values()), 'outlets': len(self.tables['outlets'])}}
